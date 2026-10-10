"""
tests/test_excel_processor.py
-----------------------------
Automated test suite verifying the hostel-to-canteen register generator.

Tests cover:
  1. 3+ students normal list in order
  2. Exactly 2 rows per student
  3. Merged Roll, Name, Room across 2 rows
  4. Unmerged 2 cells for Money/Record
  5. Preserving leading zeros (e.g. '0012')
  6. Alphanumeric room numbers (e.g. 'B-201', 'Block-C/12')
  7. Headers with different case and naming synonyms
  8. Header not on row 1 (offset headers)
  9. Error handling when required columns are missing
  10. Ensuring the original file remains unchanged
  11. Processing the actual uploaded sample files if present
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest

# Ensure project root is in sys.path
sys.path.insert(0, str(Path(__file__).parent.parent))

from excel_processor import (
    Student,
    detect_columns,
    generate_canteen_register,
    process_hostel_file,
    read_students_from_sheet,
)
from validator import validate_canteen_register


def _create_sample_excel(rows: list[list], sheet_name: str = "Sheet1") -> Path:
    """Helper to create a temporary test Excel file."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r in rows:
        ws.append(r)
    wb.save(tmp.name)
    return Path(tmp.name)


# -------------------------------------------------------------
# Core Acceptance Tests
# -------------------------------------------------------------

def test_acceptance_three_students_flow():
    """
    Acceptance test with 3 students:
    - Same students in same order
    - Each student occupies 2 rows
    - Roll, Name, Room merged across 2 rows
    - Money/record has 2 separate cells
    - No data lost
    - Original unchanged
    """
    input_data = [
        ["Roll No", "Student Name", "Room No"],
        ["240101", "Aarav Sharma", "B-101"],
        ["240102", "Diya Patel", "B-102"],
        ["240103", "Rohan Gupta", "C-205"],
    ]
    input_file = _create_sample_excel(input_data)
    output_file = input_file.parent / f"out_{input_file.name}"

    try:
        # Process
        count, info = process_hostel_file(input_file, output_file)
        assert count == 3

        # Validate
        students, _, _ = read_students_from_sheet(input_file)
        passed, issues = validate_canteen_register(output_file, students, input_file)
        assert passed, f"Validation failed with issues: {issues}"

        # Direct assertions on generated workbook
        wb_out = openpyxl.load_workbook(str(output_file))
        ws_out = wb_out.active

        # 1 header row + 3*2 = 7 total rows
        assert ws_out.max_row == 7

        # Check order and values
        expected_records = [
            ("240101", "Aarav Sharma", "B-101"),
            ("240102", "Diya Patel", "B-102"),
            ("240103", "Rohan Gupta", "C-205"),
        ]
        for i, (exp_roll, exp_name, exp_room) in enumerate(expected_records):
            r1 = 2 + i * 2
            assert str(ws_out.cell(row=r1, column=1).value) == exp_roll
            assert str(ws_out.cell(row=r1, column=2).value) == exp_name
            assert str(ws_out.cell(row=r1, column=3).value) == exp_room

        # Check merges
        merged_set = {str(r) for r in ws_out.merged_cells.ranges}
        for i in range(3):
            r1 = 2 + i * 2
            r2 = r1 + 1
            assert f"A{r1}:A{r2}" in merged_set
            assert f"B{r1}:B{r2}" in merged_set
            assert f"C{r1}:C{r2}" in merged_set
            # Record cells must NOT be in merged ranges
            assert f"D{r1}" not in str(merged_set)
            assert f"D{r2}" not in str(merged_set)

        wb_out.close()

        # Check original file remains completely unchanged
        wb_in = openpyxl.load_workbook(str(input_file))
        ws_in = wb_in.active
        assert ws_in.max_row == 4
        assert ws_in.cell(row=2, column=1).value == "240101"
        wb_in.close()

    finally:
        if input_file.exists():
            input_file.unlink()
        if output_file.exists():
            output_file.unlink()


def test_leading_zeros_and_alphanumeric_room():
    """Roll numbers with leading zeros (e.g. 0042) and complex rooms."""
    input_file = _create_sample_excel([
        ["Roll", "Name", "Hostel Room"],
        ["0042", "Kavita Rao", "Block-A/204-B"],
        ["0099", "Vikram Malhotra", "G-01"],
    ])
    output_file = input_file.parent / f"out_{input_file.name}"

    try:
        count, _ = process_hostel_file(input_file, output_file)
        assert count == 2

        wb = openpyxl.load_workbook(str(output_file))
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "0042"
        assert ws.cell(row=2, column=3).value == "Block-A/204-B"
        assert ws.cell(row=4, column=1).value == "0099"
        wb.close()
    finally:
        input_file.unlink(missing_ok=True)
        output_file.unlink(missing_ok=True)


def test_header_synonyms_and_different_order():
    """Columns in different order with alternate synonyms (Student ID, Full Name, Room/Block)."""
    input_file = _create_sample_excel([
        ["Room/Block", "Full Name", "Student ID"],
        ["101-A", "Sunil Joshi", "202499"],
    ])
    output_file = input_file.parent / f"out_{input_file.name}"

    try:
        count, info = process_hostel_file(input_file, output_file)
        assert count == 1
        assert info["roll_header"] == "Student ID"
        assert info["name_header"] == "Full Name"
        assert info["room_header"] == "Room/Block"

        wb = openpyxl.load_workbook(str(output_file))
        ws = wb.active
        # Output columns are always standard: Col A=Roll, Col B=Name, Col C=Room
        assert ws.cell(row=2, column=1).value == "202499"
        assert ws.cell(row=2, column=2).value == "Sunil Joshi"
        assert ws.cell(row=2, column=3).value == "101-A"
        wb.close()
    finally:
        input_file.unlink(missing_ok=True)
        output_file.unlink(missing_ok=True)


def test_header_not_on_row_1():
    """Notice boards or titles above the actual table headers."""
    input_file = _create_sample_excel([
        ["HOSTEL NO. 4 - MESS LIST FOR OCTOBER"],
        ["Generated on: 2026-10-01"],
        [],
        ["Roll No", "Student Name", "Room No"],
        ["240501", "Meera Sen", "F-12"],
    ])
    output_file = input_file.parent / f"out_{input_file.name}"

    try:
        count, info = process_hostel_file(input_file, output_file)
        assert count == 1
        assert info["header_row"] == 4

        wb = openpyxl.load_workbook(str(output_file))
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "240501"
        assert ws.cell(row=2, column=2).value == "Meera Sen"
        wb.close()
    finally:
        input_file.unlink(missing_ok=True)
        output_file.unlink(missing_ok=True)


def test_missing_required_column_raises_error():
    """If Name is missing, raises a clear descriptive ValueError."""
    input_file = _create_sample_excel([
        ["Roll No", "Phone Number", "Room No"],
        ["240001", "9876543210", "A-1"],
    ])
    output_file = input_file.parent / f"out_{input_file.name}"

    try:
        with pytest.raises(ValueError) as excinfo:
            process_hostel_file(input_file, output_file)
        assert "Could not automatically detect required columns" in str(excinfo.value)
    finally:
        input_file.unlink(missing_ok=True)
        output_file.unlink(missing_ok=True)


# -------------------------------------------------------------
# Sample Files Test (if available in directory)
# -------------------------------------------------------------

SAMPLE_DIR = Path(__file__).parent.parent
SAMPLE_FILES = [
    SAMPLE_DIR / "MESS LIST 2026.xlsx",
    SAMPLE_DIR / "MESS LIST 2026-v3.xlsx",
]

@pytest.mark.parametrize("sample_path", [f for f in SAMPLE_FILES if f.exists()])
def test_real_sample_files(sample_path: Path):
    """Verifies that the uploaded hostel mess lists process cleanly."""
    out_file = SAMPLE_DIR / f"test_out_{sample_path.name}"
    try:
        count, info = process_hostel_file(sample_path, out_file)
        assert count > 0, f"Expected records from {sample_path.name}"
        students, _, _ = read_students_from_sheet(sample_path)
        passed, issues = validate_canteen_register(out_file, students, sample_path)
        assert passed, f"Sample file {sample_path.name} validation failed: {issues}"
    finally:
        out_file.unlink(missing_ok=True)
