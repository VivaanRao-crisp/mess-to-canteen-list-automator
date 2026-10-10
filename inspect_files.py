"""Quick inspection script to understand the sample Excel files."""
import openpyxl

files = ["MESS LIST 2026-v3.xlsx"]

for fname in files:
    print(f"\n{'='*60}")
    print(f"FILE: {fname}")
    print('='*60)
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        for shname in wb.sheetnames:
            ws = wb[shname]
            print(f"\n  -- Sheet: '{shname}'  rows={ws.max_row}  cols={ws.max_column} --")
            for i, row in enumerate(ws.iter_rows(min_row=1,
                                                  max_row=min(30, ws.max_row),
                                                  values_only=True)):
                print(f"    Row {i+1:3d}: {row}")
            # Also print merged cell ranges
            print(f"  Merged cells: {list(ws.merged_cells.ranges)[:10]}")
    except Exception as exc:
        print(f"  ERROR: {exc}")
